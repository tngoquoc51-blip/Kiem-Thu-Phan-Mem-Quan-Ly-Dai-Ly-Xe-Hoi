"""
views/tra_luong_view.py — Tra luong nhan vien (Admin)
Hien thi bang luong, tinh luong, tra luong
"""
import os, sys, calendar
from datetime import datetime, date
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from PyQt6.QtWidgets import (
    QWidget, QVBoxLayout, QHBoxLayout, QLabel, QPushButton,
    QTableWidget, QTableWidgetItem, QHeaderView, QComboBox,
    QMessageBox, QDialog, QFormLayout, QDoubleSpinBox,
    QLineEdit, QFrame, QScrollArea
)
from PyQt6.QtCore import Qt, QTimer
from PyQt6.QtGui import QColor, QFont
from database import get_conn

STYLE = """
QWidget { font-family: 'Segoe UI', Arial; background:#ffffff; }
QTableWidget {
    background:#ffffff; alternate-background-color:#f8fafc;
    gridline-color:#e5e7eb; border:none;
    selection-background-color:#dbeafe;
}
QTableWidget::item { padding:10px 12px; color:#0f172a; font-size:14px; font-weight:600; }
QHeaderView::section {
    background:#f3f4f6; color:#0f172a;
    font-size:14px; font-weight:900; letter-spacing:0.5px;
    padding:12px; border:none; border-bottom:2px solid #2563eb;
}
QComboBox {
    background:#ffffff; color:#0f172a;
    border:1px solid #d1d5db; border-radius:8px;
    padding:8px 12px; font-size:14px; font-weight:600; min-width:140px;
}
QComboBox::drop-down { border:none; }
QComboBox QAbstractItemView {
    background:#ffffff; color:#0f172a;
    border:1px solid #d1d5db;
    selection-background-color:#e0e7ff;
}
QPushButton#btn_pay {
    background:qlineargradient(x1:0,y1:0,x2:1,y2:0,
        stop:0 #059669, stop:1 #10b981);
    color:white; border:none; border-radius:8px;
    font-size:13px; font-weight:700; padding:8px 16px;
}
QPushButton#btn_pay:hover { background:#047857; }
QPushButton#btn_pay:disabled { background:#d1fae5; color:#059669; }
QPushButton#btn_excel {
    background:#f0fdf4; color:#059669;
    border:1px solid #86efac; border-radius:8px;
    font-size:13px; font-weight:700; padding:8px 16px;
}
QPushButton#btn_excel:hover { background:#dcfce7; }
QPushButton#btn_detail {
    background:#eff6ff; color:#0284c7;
    border:1px solid #93c5fd; border-radius:8px;
    font-size:13px; font-weight:700; padding:8px 14px;
}
"""


class TraLuongView(QWidget):
    def __init__(self, current_user=None):
        super().__init__()
        self.setObjectName("page_tra_luong")
        self.setStyleSheet(STYLE)
        self.current_user = current_user or {}
        self._ensure_tra_luong_table()  # ← thêm dòng này
        self._build()
        self._load()

    def _build(self):
        root = QVBoxLayout(self)
        root.setContentsMargins(0,0,0,0); root.setSpacing(0)

        # Header
        hdr = QWidget()
        hdr.setStyleSheet("background:#ffffff;border-bottom:2px solid #e5e7eb;")
        hl = QHBoxLayout(hdr); hl.setContentsMargins(24,16,24,16); hl.setSpacing(12)
        col = QVBoxLayout(); col.setSpacing(4)
        t = QLabel("💰  Quản lý Trả lương Nhân viên")
        t.setStyleSheet("font-size:18px;font-weight:900;color:#0f172a;background:transparent;")
        s = QLabel("Tính lương • Chấm công • Thưởng phạt • Xuất bảng lương")
        s.setStyleSheet("font-size:13px;color:#6b7280;background:transparent;font-weight:600;")
        col.addWidget(t); col.addWidget(s)
        hl.addLayout(col, 1)
        self.lbl_time = QLabel()
        self.lbl_time.setStyleSheet(
            "font-size:20px;font-weight:900;color:#2563eb;"
            "background:transparent;font-family:'Courier New';")
        hl.addWidget(self.lbl_time)
        root.addWidget(hdr)

        # Toolbar
        tb = QWidget()
        tb.setStyleSheet("background:#ffffff;border-bottom:1px solid #e5e7eb;")
        tbh = QHBoxLayout(tb); tbh.setContentsMargins(20,12,20,12); tbh.setSpacing(14)

        lbl_th = QLabel("Tháng:")
        lbl_th.setStyleSheet("color:#374151;font-size:14px;background:transparent;font-weight:700;")
        self.sel_thang = QComboBox()
        for i in range(1,13): self.sel_thang.addItem(f"Tháng {i}", i)
        self.sel_thang.setCurrentIndex(datetime.now().month - 1)

        lbl_nam = QLabel("Năm:")
        lbl_nam.setStyleSheet("color:#374151;font-size:14px;background:transparent;font-weight:700;")
        self.sel_nam = QComboBox()
        for y in range(2023, 2028): self.sel_nam.addItem(str(y), y)
        self.sel_nam.setCurrentText(str(datetime.now().year))

        self.sel_thang.currentIndexChanged.connect(self._load)
        self.sel_nam.currentIndexChanged.connect(self._load)

        tbh.addWidget(lbl_th); tbh.addWidget(self.sel_thang)
        tbh.addWidget(lbl_nam); tbh.addWidget(self.sel_nam)
        tbh.addStretch()

        btn_xl = QPushButton("📊 Xuất Excel bảng lương")
        btn_xl.setObjectName("btn_excel")
        btn_xl.clicked.connect(self._export)
        tbh.addWidget(btn_xl)
        root.addWidget(tb)

        # Stats tổng
        self.stats_w = QWidget()
        self.stats_w.setStyleSheet(
            "background:#f9fafb;border-bottom:1px solid #e5e7eb;")
        sl = QHBoxLayout(self.stats_w)
        sl.setContentsMargins(18,14,18,14); sl.setSpacing(12)
        self._stat_lbls = {}
        for key, icon, label, color in [
            ("tong_nv",     "👥", "Tổng NV",      "#0284c7"),
            ("da_cham",     "✅", "Đã chấm công", "#059669"),
            ("chua_cham",   "❌", "Chưa chấm",    "#dc2626"),
            ("tong_luong",  "💰", "Tổng lương",   "#059669"),
            ("tong_thuong", "🎁", "Tổng thưởng",  "#d97706"),
            ("tong_phat",   "⚠️", "Tổng phạt",   "#ef4444"),
        ]:
            card = QWidget()
            card.setStyleSheet(
                f"background:#ffffff;border:1px solid #e5e7eb;border-radius:12px;border-top:3px solid {color};")
            cl = QVBoxLayout(card)
            cl.setContentsMargins(14,10,14,10); cl.setSpacing(2)
            lv = QLabel("—")
            lv.setStyleSheet(
                f"font-size:18px;font-weight:900;color:{color};background:transparent;")
            ll = QLabel(f"{icon} {label}")
            ll.setStyleSheet(
                "font-size:12px;color:#6b7280;background:transparent;font-weight:700;")
            cl.addWidget(lv); cl.addWidget(ll)
            sl.addWidget(card, 1)
            self._stat_lbls[key] = lv
        root.addWidget(self.stats_w)

        # Bảng lương
        cols = ["MÃ NV", "HỌ TÊN", "CHỨC VỤ", "LƯƠNG CB",
                "NGÀY CÔNG", "THỰC NHẬN CB",
                "THƯỞNG", "PHẠT", "THỰC LĨNH",
                "TRẠNG THÁI", "THAO TÁC"]
        self.tbl = QTableWidget(0, len(cols))
        self.tbl.setHorizontalHeaderLabels(cols)
        self.tbl.setAlternatingRowColors(True)
        self.tbl.setEditTriggers(QTableWidget.EditTrigger.NoEditTriggers)
        self.tbl.setShowGrid(False)
        self.tbl.verticalHeader().setVisible(False)
        self.tbl.setSelectionBehavior(QTableWidget.SelectionBehavior.SelectRows)
        h = self.tbl.horizontalHeader()
        h.setSectionResizeMode(1, QHeaderView.ResizeMode.Stretch)
        h.setSectionResizeMode(10, QHeaderView.ResizeMode.Fixed)
        self.tbl.setColumnWidth(10, 120)
        root.addWidget(self.tbl, 1)

        # Footer
        ftr = QWidget()
        ftr.setStyleSheet("background:#ffffff;border-top:1px solid #e5e7eb;")
        fl = QHBoxLayout(ftr); fl.setContentsMargins(20,12,20,12); fl.setSpacing(14)
        self.lbl_footer = QLabel("")
        self.lbl_footer.setStyleSheet("color:#374151;font-size:13px;background:transparent;font-weight:600;")
        fl.addWidget(self.lbl_footer)
        fl.addStretch()
        btn_pay_all = QPushButton("💰 Trả lương tất cả")
        btn_pay_all.setObjectName("btn_pay")
        btn_pay_all.clicked.connect(self._pay_all)
        fl.addWidget(btn_pay_all)
        root.addWidget(ftr)

        # Timer
        self._timer = QTimer(self)
        self._timer.timeout.connect(self._tick)
        self._timer.start(1000)

    def _load(self):
        thang = self.sel_thang.currentData()
        nam   = self.sel_nam.currentData()
        self.tbl.setRowCount(0)

        conn = get_conn()
        nvs  = conn.execute(
            "SELECT id, ma_nv, ho_ten, chuc_vu, luong FROM nhan_vien "
            "WHERE trang_thai != 'Nghỉ việc' ORDER BY ma_nv"
        ).fetchall()

        days_in  = calendar.monthrange(nam, thang)[1]
        work_days = sum(1 for d in range(1, days_in+1)
                       if date(nam, thang, d).weekday() < 5)

        tong_luong = tong_thuong = tong_phat = 0
        da_cham = chua_cham = 0

        for nv in nvs:
            nv_id    = nv[0]
            luong_cb = float(nv[4] or 0)

            # Chấm công
            cc_rows = conn.execute("""
                SELECT trang_thai, phat, thuong
                FROM cham_cong WHERE nv_id=?
                AND strftime('%Y-%m', ngay)=?
            """, (nv_id, f"{nam}-{thang:02d}")).fetchall()

            di_lam      = sum(1 for r in cc_rows if r[0] != "Vắng mặt")
            tong_p      = sum(float(r[1] or 0) for r in cc_rows)
            tong_t      = sum(float(r[2] or 0) for r in cc_rows)
            vang        = sum(1 for r in cc_rows if r[0] == "Vắng mặt")
            muon        = sum(1 for r in cc_rows if r[0] == "Đi muộn")

            if work_days > 0:
                luong_thuc_cb = (luong_cb / work_days) * di_lam
            else:
                luong_thuc_cb = luong_cb

            # Thưởng chuyên cần
            if vang == 0 and muon == 0 and di_lam > 0:
                tong_t += 200_000

            thuc_linh = luong_thuc_cb - tong_p + tong_t

            # Trạng thái trả lương
            try:
                tt_luong = conn.execute("""
                            SELECT trang_thai FROM tra_luong
                            WHERE nv_id=? AND thang=? AND nam=?
                        """, (nv_id, thang, nam)).fetchone()
                tt_str = tt_luong[0] if tt_luong else "Chưa trả"
            except:
                tt_str = "Chưa trả"

            if di_lam > 0: da_cham += 1
            else: chua_cham += 1

            tong_luong += thuc_linh
            tong_thuong += tong_t
            tong_phat   += tong_p

            # Render row
            r = self.tbl.rowCount()
            self.tbl.insertRow(r)
            self.tbl.setRowHeight(r, 46)

            tt_color = "#059669" if tt_str == "Đã trả" else "#d97706"

            data = [
                (nv[1], "#7c3aed"),
                (nv[2], "#0f172a"),
                (nv[3], "#475569"),
                (f"{luong_cb/1e6:.1f}tr", "#7c3aed"),
                (f"{di_lam}/{work_days}", "#0284c7"),
                (f"{luong_thuc_cb/1e6:.2f}tr", "#0f172a"),
                (f"+{tong_t/1e3:.0f}k" if tong_t > 0 else "—", "#059669"),
                (f"-{tong_p/1e3:.0f}k" if tong_p > 0 else "—", "#dc2626"),
                (f"{thuc_linh/1e6:.2f}tr", "#059669"),
                (tt_str, tt_color),
            ]

            for c, (val, color) in enumerate(data):
                item = QTableWidgetItem(val)
                item.setData(Qt.ItemDataRole.UserRole, nv_id)
                item.setForeground(QColor(color))
                if c in [3, 8]:
                    item.setFont(QFont("Segoe UI", 14, QFont.Weight.Bold))
                self.tbl.setItem(r, c, item)

            # Nút trả lương
            btn = QPushButton(
                "✅ Đã trả" if tt_str == "Đã trả" else "💳 Trả lương")
            btn.setObjectName("btn_pay")
            if tt_str == "Đã trả":
                btn.setEnabled(False)
                btn.setStyleSheet(
                    "background:#d1fae5;color:#059669;border:none;"
                    "border-radius:8px;font-size:12px;font-weight:700;padding:7px 12px;")
            else:
                btn.setStyleSheet(
                    "background:#059669;color:white;border:none;"
                    "border-radius:8px;font-size:12px;font-weight:700;padding:7px 12px;")
                btn.clicked.connect(
                    lambda _, nid=nv_id, ten=nv[2], tl=thuc_linh:
                    self._pay_one(nid, ten, tl))
            self.tbl.setCellWidget(r, 10, btn)

        conn.close()

        # Stats
        self._stat_lbls["tong_nv"].setText(str(len(nvs)))
        self._stat_lbls["da_cham"].setText(str(da_cham))
        self._stat_lbls["chua_cham"].setText(str(chua_cham))
        self._stat_lbls["tong_luong"].setText(f"{tong_luong/1e6:.1f}tr")
        self._stat_lbls["tong_thuong"].setText(f"{tong_thuong/1e3:.0f}k")
        self._stat_lbls["tong_phat"].setText(f"{tong_phat/1e3:.0f}k")
        self.lbl_footer.setText(
            f"Tháng {thang}/{nam} • {len(nvs)} nhân viên • "
            f"Tổng quỹ lương: {tong_luong/1e6:.2f} triệu ₫")

    def _ensure_tra_luong_table(self):
        conn = get_conn()
        conn.execute("""CREATE TABLE IF NOT EXISTS tra_luong (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            nv_id INTEGER NOT NULL,
            thang INTEGER NOT NULL,
            nam INTEGER NOT NULL,
            so_tien REAL DEFAULT 0,
            trang_thai TEXT DEFAULT 'Chua tra',
            ngay_tra TEXT,
            ghi_chu TEXT,
            created_at TEXT DEFAULT CURRENT_TIMESTAMP
        )""")
        conn.commit(); conn.close()

    def _pay_one(self, nv_id, ho_ten, so_tien):
        thang = self.sel_thang.currentData()
        nam   = self.sel_nam.currentData()

        reply = QMessageBox.question(self,
            "Xác nhận trả lương",
            f"Trả lương cho {ho_ten}?\n"
            f"Số tiền: {so_tien/1e6:.2f} triệu ₫\n"
            f"Tháng {thang}/{nam}",
            QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No)

        if reply == QMessageBox.StandardButton.Yes:
            self._ensure_tra_luong_table()
            conn = get_conn()
            try:
                conn.execute("""
                    INSERT OR REPLACE INTO tra_luong
                    (nv_id, thang, nam, so_tien, trang_thai, ngay_tra)
                    VALUES (?,?,?,?,'Đã trả', date('now'))
                """, (nv_id, thang, nam, so_tien))
                conn.commit()
                QMessageBox.information(self, "✅ Thành công",
                    f"Đã trả lương cho {ho_ten}!\n"
                    f"Số tiền: {so_tien/1e6:.2f} triệu ₫")
                self._load()
            except Exception as e:
                QMessageBox.critical(self, "Lỗi", str(e))
            finally:
                conn.close()

    def _pay_all(self):
        thang = self.sel_thang.currentData()
        nam   = self.sel_nam.currentData()

        reply = QMessageBox.question(self,
            "Trả lương tất cả",
            f"Xác nhận trả lương TẤT CẢ nhân viên\n"
            f"Tháng {thang}/{nam}?",
            QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No)

        if reply == QMessageBox.StandardButton.Yes:
            self._ensure_tra_luong_table()
            conn = get_conn()
            try:
                nvs = conn.execute(
                    "SELECT id FROM nhan_vien WHERE trang_thai != 'Nghỉ việc'"
                ).fetchall()
                for nv in nvs:
                    conn.execute("""
                        INSERT OR IGNORE INTO tra_luong
                        (nv_id, thang, nam, trang_thai, ngay_tra)
                        VALUES (?,?,?,'Đã trả', date('now'))
                    """, (nv[0], thang, nam))
                conn.commit()
                QMessageBox.information(self, "✅ Thành công",
                    f"Đã trả lương tất cả {len(nvs)} nhân viên!")
                self._load()
            except Exception as e:
                QMessageBox.critical(self, "Lỗi", str(e))
            finally:
                conn.close()

    def _export(self):
        try:
            import openpyxl
            from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
            from openpyxl.utils import get_column_letter
            import datetime as dt

            thang = self.sel_thang.currentData()
            nam   = self.sel_nam.currentData()

            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = f"Bang luong T{thang}-{nam}"

            # ── Tiêu đề lớn ──────────────────────────────
            ws.merge_cells("A1:J1")
            ws["A1"] = f"BẢNG LƯƠNG THÁNG {thang}/{nam} — AUTOVIET"
            ws["A1"].font = Font(name="Segoe UI", size=14, bold=True, color="FFFFFF")
            ws["A1"].fill = PatternFill("solid", fgColor="1E3A5F")
            ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
            ws.row_dimensions[1].height = 38

            # ── Ngày xuất ────────────────────────────────
            ws.merge_cells("A2:J2")
            ws["A2"] = f"Ngày xuất: {dt.datetime.now().strftime('%d/%m/%Y %H:%M')}"
            ws["A2"].font = Font(name="Segoe UI", size=10, italic=True, color="64748B")
            ws["A2"].alignment = Alignment(horizontal="right")
            ws.row_dimensions[2].height = 20

            # ── Header ───────────────────────────────────
            headers = ["MÃ NV", "HỌ TÊN", "CHỨC VỤ", "LƯƠNG CB",
                       "NGÀY CÔNG", "THỰC NHẬN CB", "THƯỞNG", "PHẠT",
                       "THỰC LĨNH", "TRẠNG THÁI"]
            col_widths = [10, 22, 18, 14, 12, 16, 12, 12, 14, 12]
            thin = Side(style="thin", color="BFDBFE")
            border = Border(left=thin, right=thin, top=thin, bottom=thin)

            for i, (h, w) in enumerate(zip(headers, col_widths), 1):
                cell = ws.cell(row=3, column=i, value=h)
                cell.font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="1E40AF")
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.border = border
                ws.column_dimensions[get_column_letter(i)].width = w
            ws.row_dimensions[3].height = 28

            # ── Dữ liệu ──────────────────────────────────
            fill_white = PatternFill("solid", fgColor="FFFFFF")
            fill_alt   = PatternFill("solid", fgColor="EFF6FF")
            fill_paid  = PatternFill("solid", fgColor="F0FDF4")
            fill_neg   = PatternFill("solid", fgColor="FFF1F2")

            for r in range(self.tbl.rowCount()):
                ri = r + 4
                ws.row_dimensions[ri].height = 24

                # Lấy giá trị trạng thái để chọn màu nền
                tt_item = self.tbl.item(r, 9)
                tt = tt_item.text() if tt_item else ""

                # Lấy thực lĩnh để check âm
                tl_item = self.tbl.item(r, 8)
                tl_text = tl_item.text() if tl_item else "0"
                is_neg = tl_text.startswith("-")

                if tt == "Đã trả":
                    row_fill = fill_paid
                elif is_neg:
                    row_fill = fill_neg
                elif r % 2 == 0:
                    row_fill = fill_white
                else:
                    row_fill = fill_alt

                for c in range(10):
                    item = self.tbl.item(r, c)
                    val = item.text() if item else ""
                    cell = ws.cell(row=ri, column=c+1, value=val)
                    cell.fill = row_fill
                    cell.border = border
                    cell.alignment = Alignment(vertical="center",
                        horizontal="center" if c in (0, 4, 9) else "left")

                    # Màu chữ theo cột
                    if c == 0:
                        cell.font = Font(name="Segoe UI", size=11, bold=True, color="2563EB")
                    elif c == 3:
                        cell.font = Font(name="Segoe UI", size=11, bold=True, color="0F172A")
                    elif c == 6:  # Thưởng
                        cell.font = Font(name="Segoe UI", size=11, bold=True, color="059669")
                    elif c == 7:  # Phạt
                        cell.font = Font(name="Segoe UI", size=11, bold=True, color="DC2626")
                    elif c == 8:  # Thực lĩnh
                        color = "DC2626" if is_neg else "059669"
                        cell.font = Font(name="Segoe UI", size=12, bold=True, color=color)
                    elif c == 9:  # Trạng thái
                        color = "059669" if tt == "Đã trả" else "D97706"
                        cell.font = Font(name="Segoe UI", size=11, bold=True, color=color)
                    else:
                        cell.font = Font(name="Segoe UI", size=11, color="374151")

            # ── Tổng kết ─────────────────────────────────
            last_row = self.tbl.rowCount() + 4
            ws.row_dimensions[last_row].height = 28
            ws.merge_cells(f"A{last_row}:H{last_row}")
            ws[f"A{last_row}"] = "TỔNG CỘNG"
            ws[f"A{last_row}"].font = Font(name="Segoe UI", size=12, bold=True, color="FFFFFF")
            ws[f"A{last_row}"].fill = PatternFill("solid", fgColor="1E3A5F")
            ws[f"A{last_row}"].alignment = Alignment(horizontal="right", vertical="center")
            ws[f"A{last_row}"].border = border

            # Tính tổng thực lĩnh
            tong_tl = self.lbl_tong_luong.text() if hasattr(self, 'lbl_tong_luong') else ""
            ws[f"I{last_row}"] = tong_tl
            ws[f"I{last_row}"].font = Font(name="Segoe UI", size=12, bold=True, color="FFFFFF")
            ws[f"I{last_row}"].fill = PatternFill("solid", fgColor="1E3A5F")
            ws[f"I{last_row}"].alignment = Alignment(horizontal="center", vertical="center")
            ws[f"I{last_row}"].border = border
            ws[f"J{last_row}"].fill = PatternFill("solid", fgColor="1E3A5F")
            ws[f"J{last_row}"].border = border

            ws.freeze_panes = "A4"
            fname = f"BangLuong_T{thang}_{nam}_{dt.datetime.now().strftime('%d%m%Y_%H%M')}.xlsx"
            wb.save(fname)
            QMessageBox.information(self, "✅ Xuất thành công!", f"Đã xuất: {fname}")

        except Exception as e:
            QMessageBox.critical(self, "Lỗi", str(e))

    def refresh(self):
        self._load()

    def _tick(self):
        self.lbl_time.setText(datetime.now().strftime("%H:%M:%S"))