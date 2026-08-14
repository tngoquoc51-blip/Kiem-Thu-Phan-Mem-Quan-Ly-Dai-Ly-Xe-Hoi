"""
views/lich_baoduong_view.py — Lịch bảo dưỡng xe định kỳ
File MỚI — thêm vào views/
"""
import sys, os
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from PyQt6.QtWidgets import (
    QWidget, QVBoxLayout, QHBoxLayout, QLabel, QPushButton,
    QTableWidget, QTableWidgetItem, QHeaderView, QFrame,
    QDialog, QFormLayout, QComboBox, QLineEdit, QMessageBox,
    QDoubleSpinBox, QScrollArea, QGridLayout
)
from PyQt6.QtCore import Qt, QDate
from PyQt6.QtGui import QColor, QFont
import pandas as pd
from database import get_conn
from datetime import datetime, date, timedelta


STYLE = """
QWidget { font-family: 'Segoe UI', Arial; }
QWidget#toolbar_widget {
    background: #ffffff;
    border-bottom: 1px solid #e5e7eb;
}
QPushButton#btn_add {
    background: qlineargradient(x1:0,y1:0,x2:1,y2:0,
        stop:0 #6d28d9, stop:1 #7c3aed);
    color:white; border:none; border-radius:9px;
    font-size:13px; font-weight:700; padding:9px 18px;
}
QPushButton#btn_add:hover { background:#7c3aed; }
QPushButton#btn_del {
    background:#450a0a; color:#fca5a5;
    border:1px solid #7f1d1d; border-radius:8px;
    padding:8px 14px; font-size:12px;
}
QPushButton#btn_del:hover { background:#7f1d1d; }
QPushButton#btn_done {
    background:#052e16; color:#86efac;
    border:1px solid #166534; border-radius:8px;
    padding:8px 14px; font-size:12px; font-weight:600;
}
QPushButton#btn_done:hover { background:#15803d; }
QPushButton#btn_excel {
    background:#14532d; color:#86efac;
    border:1px solid #166534; border-radius:8px;
    padding:8px 14px; font-size:12px;
}
QLineEdit#search_box {
    background:#f9fafb; color:#111827;
    border:1px solid #e5e7eb; border-radius:9px;
    padding:8px 16px; font-size:13px; min-width:240px;
}
QLineEdit#search_box:focus { border-color:#2563eb; background:#ffffff; }
QTableWidget {
    background:#ffffff; alternate-background-color:#f9fafb;
    gridline-color:#f1f5f9; border:none;
}
QTableWidget::item { padding:8px 12px; color:#1e293b; }
QHeaderView::section {
    background:#ffffff; color:#2563eb; font-size:12px;
    font-weight:900; letter-spacing:1px; padding:14px 12px; border:none;
    border-bottom: 2px solid #2563eb;
}
QDialog { background:#ffffff; }
QLabel#dlg_lbl { color:#6b7280; font-size:11px; font-weight:700;
    background:transparent; text-transform:uppercase; letter-spacing:1px; }
QLineEdit, QComboBox, QDoubleSpinBox {
    background:#f9fafb; color:#111827;
    border:1px solid #e5e7eb; border-radius:8px;
    padding:9px 12px; font-size:13px;
}
QLineEdit:focus, QComboBox:focus { border-color:#2563eb; background:#ffffff; }
QPushButton#dlg_save {
    background:#2563eb; color:white; border:none;
    border-radius:9px; font-size:14px; font-weight:700;
    padding:11px 24px;
}
QPushButton#dlg_save:hover { background:#1d4ed8; }
QPushButton#dlg_cancel {
    background:#f3f4f6; color:#6b7280;
    border:1px solid #d1d5db; border-radius:9px;
    font-size:13px; padding:10px 20px;
}
"""


def _init_lich_table():
    """Tạo bảng lich_baoduong nếu chưa có"""
    conn = get_conn()
    conn.execute("""
        CREATE TABLE IF NOT EXISTS lich_baoduong (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            xe_id INTEGER REFERENCES xe(id),
            kh_id INTEGER REFERENCES khach_hang(id),
            loai TEXT DEFAULT 'Bảo dưỡng định kỳ',
            ngay_hen TEXT,
            km_hien_tai INTEGER DEFAULT 0,
            ghi_chu TEXT,
            trang_thai TEXT DEFAULT 'Chờ xác nhận',
            created_at TEXT DEFAULT (date('now'))
        )
    """)
    conn.commit(); conn.close()


class LichBaoDuongView(QWidget):
    def __init__(self, current_user=None):
        super().__init__()
        self.setObjectName("page_lich_baoduong")
        self.setStyleSheet(STYLE)
        self.current_user = current_user or {}
        self._rows = []; self._sel_id = None
        _init_lich_table()
        self._build()
        self._load()

    def _build(self):
        root = QVBoxLayout(self); root.setContentsMargins(0,0,0,0); root.setSpacing(0)

        # Toolbar
        tb = QWidget(); tb.setObjectName("toolbar_widget")
        tbh = QHBoxLayout(tb); tbh.setContentsMargins(16,10,16,10); tbh.setSpacing(8)

        title = QLabel("📅  Lịch bảo dưỡng xe định kỳ")
        title.setStyleSheet("font-size:17px;font-weight:900;color:#111827;background:transparent;")

        btn_add  = QPushButton("➕  Đặt lịch mới"); btn_add.setObjectName("btn_add")
        btn_done = QPushButton("✅  Đã hoàn thành"); btn_done.setObjectName("btn_done")
        btn_del  = QPushButton("🗑  Xoá"); btn_del.setObjectName("btn_del")
        btn_xl   = QPushButton("📊 Excel"); btn_xl.setObjectName("btn_excel")

        for b in [btn_add,btn_done,btn_del,btn_xl]:
            b.setCursor(Qt.CursorShape.PointingHandCursor)

        self.search = QLineEdit(); self.search.setObjectName("search_box")
        self.search.setPlaceholderText("🔍  Tìm theo xe, khách hàng...")
        self.search.textChanged.connect(lambda t: self._load(t.strip()))

        tbh.addWidget(title); tbh.addWidget(btn_add); tbh.addWidget(btn_done)
        tbh.addWidget(btn_del); tbh.addWidget(btn_xl); tbh.addStretch()
        tbh.addWidget(self.search)
        root.addWidget(tb)

        # Stat cards
        stat_w = QWidget();
        stat_w.setStyleSheet("background:#f9fafb;padding:14px 16px 12px;border-bottom:1px solid #e5e7eb;")
        stat_h = QHBoxLayout(stat_w);
        stat_h.setSpacing(12)
        self.sc_cho = self._stat("⏳", "Chờ xác nhận", "0", "#fbbf24")
        self.sc_sap = self._stat("📅", "Sắp đến hạn", "0", "#f97316")
        self.sc_tre = self._stat("⚠️", "Quá hạn", "0", "#ef4444")
        self.sc_hoan = self._stat("✅", "Hoàn thành", "0", "#10b981")
        for sc in [self.sc_cho,self.sc_sap,self.sc_tre,self.sc_hoan]:
            stat_h.addWidget(sc)
        root.addWidget(stat_w)

        # Table
        cols = ["ID", "XE", "KHÁCH HÀNG", "LOẠI BẢO DƯỠNG",
                "NGÀY HẸN", "SỐ KM", "GHI CHÚ", "TRẠNG THÁI"]
        self.tbl = QTableWidget(0, len(cols))
        self.tbl.setHorizontalHeaderLabels(cols)
        self.tbl.setAlternatingRowColors(True)
        self.tbl.setSelectionBehavior(QTableWidget.SelectionBehavior.SelectRows)
        self.tbl.setEditTriggers(QTableWidget.EditTrigger.NoEditTriggers)
        self.tbl.setShowGrid(False)
        self.tbl.verticalHeader().setVisible(False)
        h = self.tbl.horizontalHeader()
        h.hideSection(0)  # ẩn cột ID
        h.setSectionResizeMode(1, QHeaderView.ResizeMode.Stretch)  # XE - rộng
        h.setSectionResizeMode(2, QHeaderView.ResizeMode.ResizeToContents)  # KHÁCH HÀNG
        h.setSectionResizeMode(3, QHeaderView.ResizeMode.ResizeToContents)  # LOẠI BẢO DƯỠNG
        h.setSectionResizeMode(4, QHeaderView.ResizeMode.ResizeToContents)  # NGÀY HẸN
        h.setSectionResizeMode(5, QHeaderView.ResizeMode.ResizeToContents)  # SỐ KM
        h.setSectionResizeMode(6, QHeaderView.ResizeMode.Stretch)  # GHI CHÚ - rộng
        h.setSectionResizeMode(7, QHeaderView.ResizeMode.ResizeToContents)  # TRẠNG THÁI
        self.tbl.selectionModel().selectionChanged.connect(self._on_sel)
        root.addWidget(self.tbl)

        btn_add.clicked.connect(self._add)
        btn_done.clicked.connect(self._mark_done)
        btn_del.clicked.connect(self._delete)
        btn_xl.clicked.connect(self._export)

    def _stat(self, icon, label, val, color):
        w = QWidget()
        w.setStyleSheet(f"background:#ffffff;border-radius:10px;border:1px solid #e5e7eb;border-top:4px solid {color};")
        lv = QVBoxLayout(w);
        lv.setContentsMargins(16, 12, 16, 12);
        lv.setSpacing(6)

        # Icon
        icon_lbl = QLabel(icon)
        icon_lbl.setStyleSheet(f"font-size:28px;background:transparent;")
        icon_lbl.setAlignment(Qt.AlignmentFlag.AlignCenter)
        lv.addWidget(icon_lbl)

        # Label text
        label_lbl = QLabel(label)
        label_lbl.setStyleSheet(
            f"font-size:13px;color:{color};font-weight:800;letter-spacing:.8px;background:transparent;text-align:center;")
        label_lbl.setAlignment(Qt.AlignmentFlag.AlignCenter)
        lv.addWidget(label_lbl)

        # Value number
        vl = QLabel(val)
        vl.setStyleSheet(f"font-size:28px;font-weight:900;color:{color};background:transparent;")
        vl.setAlignment(Qt.AlignmentFlag.AlignCenter)
        lv.addWidget(vl)

        w._val = vl
        return w

    def _lbl(self, text, style=""):
        l = QLabel(text); l.setStyleSheet(style+"background:transparent;"); return l

    def _load(self, q=""):
        conn = get_conn()
        sql = """SELECT lb.*, x.hang_xe||' '||x.dong_xe as ten_xe, kh.ho_ten as ten_kh
                 FROM lich_baoduong lb
                 LEFT JOIN xe x ON lb.xe_id=x.id
                 LEFT JOIN khach_hang kh ON lb.kh_id=kh.id"""
        p = []
        if q:
            sql += " WHERE (x.hang_xe LIKE ? OR x.dong_xe LIKE ? OR kh.ho_ten LIKE ?)"
            p = [f"%{q}%"]*3
        sql += " ORDER BY lb.ngay_hen ASC"
        self._rows = [dict(r) for r in conn.execute(sql, p).fetchall()]
        conn.close()

        today = date.today()
        soon = today + timedelta(days=7)

        cnt_cho = cnt_sap = cnt_tre = cnt_hoan = 0
        STATUS_COL = {
            "Chờ xác nhận":"#fbbf24",
            "Đã xác nhận":"#60a5fa",
            "Hoàn thành":"#4ade80",
            "Huỷ":"#f87171",
        }

        self.tbl.setRowCount(0)
        for row in self._rows:
            r = self.tbl.rowCount(); self.tbl.insertRow(r)
            self.tbl.setRowHeight(r, 48)

            ngay_hen = row.get("ngay_hen","") or ""
            tt = row.get("trang_thai","")

            # Xác định trạng thái thực
            display_tt = tt
            row_color = None
            if tt not in ["Hoàn thành","Huỷ"] and ngay_hen:
                try:
                    ngay = date.fromisoformat(ngay_hen)
                    if ngay < today:
                        display_tt = "⚠️ Quá hạn"; row_color = "#2a0a0a"; cnt_tre += 1
                    elif ngay <= soon:
                        display_tt = "📅 Sắp đến"; row_color = "#1c1200"; cnt_sap += 1
                    else:
                        cnt_cho += 1
                except: cnt_cho += 1
            elif tt == "Hoàn thành": cnt_hoan += 1
            else: cnt_cho += 1

            vals = [str(row.get("id","")),
                    row.get("ten_xe","") or "—",
                    row.get("ten_kh","") or "—",
                    row.get("loai",""),
                    ngay_hen,
                    f"{int(row.get('km_hien_tai',0) or 0):,} km",
                    row.get("ghi_chu","") or "—",
                    display_tt]

            for c, val in enumerate(vals):
                item = QTableWidgetItem(val)
                item.setData(Qt.ItemDataRole.UserRole, row.get("id"))

                # ✅ Font chung cho tất cả
                item.setFont(QFont("Segoe UI", 13, QFont.Weight.Bold))
                item.setForeground(QColor("#1e293b"))

                # ✅ Column 7 = display_tt (trạng thái)
                if c == 7:
                    col = STATUS_COL.get(tt, "#fbbf24")
                    if "Quá hạn" in display_tt:
                        col = "#dc2626"
                    elif "Sắp đến" in display_tt:
                        col = "#f97316"
                    item.setForeground(QColor(col))
                    item.setFont(QFont("Segoe UI", 13, QFont.Weight.Bold))

                if row_color and c > 0:
                    item.setBackground(QColor(row_color))

                self.tbl.setItem(r, c, item)

        self.sc_cho._val.setText(str(cnt_cho))
        self.sc_sap._val.setText(str(cnt_sap))
        self.sc_tre._val.setText(str(cnt_tre))
        self.sc_hoan._val.setText(str(cnt_hoan))

    def _on_sel(self):
        r = self.tbl.currentRow()
        if 0 <= r < len(self._rows):
            self._sel_id = self._rows[r]["id"]

    def _add(self):
        if DatLichDialog(self).exec(): self._load()

    def _mark_done(self):
        if not self._sel_id:
            QMessageBox.warning(self,"","Chọn lịch cần đánh dấu hoàn thành!"); return
        conn = get_conn()
        conn.execute("UPDATE lich_baoduong SET trang_thai='Hoàn thành' WHERE id=?", (self._sel_id,))
        conn.commit(); conn.close()
        self._load()

    def _delete(self):
        if not self._sel_id:
            QMessageBox.warning(self,"","Chọn lịch cần xoá!"); return
        if QMessageBox.question(self,"Xác nhận","Xoá lịch bảo dưỡng này?",
            QMessageBox.StandardButton.Yes|QMessageBox.StandardButton.No
        )==QMessageBox.StandardButton.Yes:
            conn=get_conn(); conn.execute("DELETE FROM lich_baoduong WHERE id=?",(self._sel_id,))
            conn.commit(); conn.close(); self._sel_id=None; self._load()

    def _export(self):
        import openpyxl
        from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
        import datetime

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Lich bao duong"

        ws.merge_cells("A1:I1")
        ws["A1"] = "LỊCH BẢO DƯỠNG XE — AUTOVIET"
        ws["A1"].font = Font(name="Segoe UI", size=14, bold=True, color="FFFFFF")
        ws["A1"].fill = PatternFill("solid", fgColor="0F766E")
        ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 36

        ws.merge_cells("A2:I2")
        ws["A2"] = f"Ngày xuất: {datetime.datetime.now().strftime('%d/%m/%Y %H:%M')}"
        ws["A2"].font = Font(name="Segoe UI", size=10, italic=True, color="64748B")
        ws["A2"].alignment = Alignment(horizontal="right")
        ws.row_dimensions[2].height = 20

        headers = ["STT", "XE", "KHÁCH HÀNG", "LOẠI BD", "NGÀY HẸN",
                   "KM HIỆN TẠI", "GHI CHÚ", "TRẠNG THÁI", "NGÀY TẠO"]
        col_widths = [6, 28, 22, 18, 14, 14, 25, 16, 14]
        thin = Side(style="thin", color="99F6E4")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)

        for i, (h, w) in enumerate(zip(headers, col_widths), 1):
            cell = ws.cell(row=3, column=i, value=h)
            cell.font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="0F766E")
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = border
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.row_dimensions[3].height = 28

        fill_white = PatternFill("solid", fgColor="FFFFFF")
        fill_alt   = PatternFill("solid", fgColor="F0FDFA")

        STATUS_COLOR = {
            "Hoàn thành":     "059669",
            "Đã nhắc":        "2563EB",
            "Chờ xử lý":      "D97706",
            "Sắp đến hạn":    "EA580C",
            "Quá hạn":        "DC2626",
        }

        for ri, row in enumerate(self._rows, 4):
            fill = fill_white if ri % 2 == 0 else fill_alt
            ws.row_dimensions[ri].height = 24

            tt = str(row.get("trang_thai", "") or "")

            vals = [
                ri - 3,
                row.get("ten_xe", "") or "—",
                row.get("ten_kh", "") or "—",
                row.get("loai", "") or "—",
                row.get("ngay_hen", "") or "—",
                row.get("km_hien", "") or "—",
                row.get("ghi_chu", "") or "—",
                tt,
                row.get("created_at", "") or "—",
            ]

            for ci, val in enumerate(vals, 1):
                cell = ws.cell(row=ri, column=ci, value=val)
                cell.fill = fill
                cell.border = border
                cell.alignment = Alignment(vertical="center",
                    horizontal="center" if ci in (1, 5, 6, 8, 9) else "left")

                if ci == 1:
                    cell.font = Font(name="Segoe UI", size=11, color="64748B")
                elif ci == 2:
                    cell.font = Font(name="Segoe UI", size=11, bold=True, color="0F172A")
                elif ci == 3:
                    cell.font = Font(name="Segoe UI", size=11, color="0F172A")
                elif ci == 5:
                    cell.font = Font(name="Segoe UI", size=11, bold=True, color="0F766E")
                elif ci == 8:
                    color = STATUS_COLOR.get(tt, "64748B")
                    cell.font = Font(name="Segoe UI", size=11, bold=True, color=color)
                else:
                    cell.font = Font(name="Segoe UI", size=11, color="374151")

        ws.freeze_panes = "A4"
        fname = f"LichBaoDuong_{datetime.datetime.now().strftime('%d%m%Y_%H%M')}.xlsx"
        wb.save(fname)
        QMessageBox.information(self, "✅ Xuất thành công!", f"Đã xuất: {fname}")

    def refresh(self): self._load(self.search.text())


class DatLichDialog(QDialog):
    def __init__(self, parent=None):
        super().__init__(parent)
        self.setWindowTitle("Đặt lịch bảo dưỡng")
        self.setMinimumWidth(480)
        self.setStyleSheet(STYLE)
        self._build()

    def _build(self):
        outer = QVBoxLayout(self); outer.setContentsMargins(24,20,24,20); outer.setSpacing(14)

        title = QLabel("📅  ĐẶT LỊCH BẢO DƯỠNG XE")
        title.setStyleSheet("font-size:16px;font-weight:900;color:#111827;background:transparent;")
        sep = QFrame(); sep.setFrameShape(QFrame.Shape.HLine)
        sep.setStyleSheet("background:#252840;max-height:1px;")
        outer.addWidget(title); outer.addWidget(sep)

        form = QFormLayout(); form.setSpacing(10)
        form.setLabelAlignment(Qt.AlignmentFlag.AlignRight)

        def lbl(t, color="#1d4ed8"):
            l = QLabel(t)
            l.setObjectName("dlg_lbl")
            l.setStyleSheet(f"color:{color};font-size:14px;font-weight:900;background:transparent;letter-spacing:1px;")
            return l

        conn = get_conn()
        xe_rows = conn.execute("SELECT id,ma_xe,hang_xe,dong_xe FROM xe ORDER BY id DESC").fetchall()
        kh_rows = conn.execute("SELECT id,ma_kh,ho_ten FROM khach_hang ORDER BY id DESC").fetchall()
        conn.close()

        self.f_xe = QComboBox()
        for r in xe_rows: self.f_xe.addItem(f"{r[1]} — {r[2]} {r[3]}", r[0])
        self.f_kh = QComboBox()
        for r in kh_rows: self.f_kh.addItem(f"{r[1]} — {r[2]}", r[0])
        self.f_loai = QComboBox()
        self.f_loai.addItems(["Bảo dưỡng định kỳ","Thay dầu & lọc","Thay lốp xe",
                               "Kiểm tra tổng quát","Bảo dưỡng điều hòa","Đăng kiểm","Khác"])
        self.f_ngay = QLineEdit(); self.f_ngay.setPlaceholderText("YYYY-MM-DD")
        # Điền ngày mặc định 7 ngày tới
        from datetime import date, timedelta
        self.f_ngay.setText((date.today()+timedelta(days=7)).strftime("%Y-%m-%d"))
        self.f_km = QLineEdit("0"); self.f_km.setPlaceholderText("Số km hiện tại")
        self.f_ghi = QLineEdit(); self.f_ghi.setPlaceholderText("Ghi chú thêm...")

        colors = ["#2563eb", "#16a34a", "#dc2626", "#f97316", "#7c3aed", "#0891b2"]
        for (txt, w), color in zip([("Xe *", self.f_xe), ("Khách hàng", self.f_kh),
                                    ("Loại bảo dưỡng *", self.f_loai), ("Ngày hẹn *", self.f_ngay),
                                    ("Số km hiện tại", self.f_km), ("Ghi chú", self.f_ghi)], colors):
            form.addRow(lbl(txt, color), w)
        outer.addLayout(form)

        bh = QHBoxLayout(); bh.addStretch()
        btn_c = QPushButton("Huỷ bỏ"); btn_c.setObjectName("dlg_cancel")
        btn_c.clicked.connect(self.reject)
        btn_s = QPushButton("📅  Đặt lịch"); btn_s.setObjectName("dlg_save")
        btn_s.clicked.connect(self._save); btn_s.setDefault(True)
        bh.addWidget(btn_c); bh.addWidget(btn_s)
        outer.addLayout(bh)

    def _save(self):
        ngay = self.f_ngay.text().strip()
        if not ngay:
            QMessageBox.warning(self,"","Chọn ngày hẹn!"); return
        try:
            date.fromisoformat(ngay)
        except:
            QMessageBox.warning(self,"","Ngày không hợp lệ! Dùng định dạng YYYY-MM-DD"); return

        conn = get_conn()
        try:
            conn.execute("""
                INSERT INTO lich_baoduong(xe_id,kh_id,loai,ngay_hen,km_hien_tai,ghi_chu)
                VALUES(?,?,?,?,?,?)
            """, (self.f_xe.currentData(), self.f_kh.currentData(),
                  self.f_loai.currentText(), ngay,
                  int(self.f_km.text() or 0),
                  self.f_ghi.text()))
            conn.commit()
            QMessageBox.information(self,"✅ OK","Đặt lịch bảo dưỡng thành công!")
            self.accept()
        except Exception as e: QMessageBox.critical(self,"Lỗi",str(e))
        finally: conn.close()
